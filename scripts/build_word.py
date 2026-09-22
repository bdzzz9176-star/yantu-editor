from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "dalian_maritime" / "draft.md"
CHART = ROOT / "outputs" / "dalian_maritime" / "charts" / "cutoff_trends.png"
OUTPUT = ROOT / "outputs" / "dalian_maritime" / "大连海事大学电子通信考研择校分析_V1.docx"
GREEN = "18563F"
GREEN_SOFT = "E8F0EB"
GOLD = "B66A20"
GOLD_SOFT = "FFF1DC"
INK = "1A211E"
MUTED = "68736F"
LINE = "D8DDD9"


def set_font(run, size=10.5, bold=False, color=INK, name="宋体"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
        if node.getparent() is None: mar.append(node)


def add_rich_text(paragraph, text, size=10.5, color=INK):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part: continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        set_font(run, size=size, bold=bold, color=color)


def style_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
    section.left_margin, section.right_margin = Cm(2.25), Cm(2.25)
    section.header_distance, section.footer_distance = Cm(1.1), Cm(1.1)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.space_after = Pt(3)
    for style_name, size, color, before, after in (
        ("Title", 22, INK, 0, 12), ("Heading 1", 15, GREEN, 14, 6),
        ("Heading 2", 12.5, INK, 10, 5), ("Heading 3", 11, GOLD, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_identity_band(doc):
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(2.2), Cm(5.2), Cm(2.2), Cm(5.2)]
    values = [("学校", "大连海事大学", "学院", "信息科学技术学院"),
              ("层次", "211 · 双一流", "学科", "信息与通信工程 B-")]
    for row, row_values in zip(table.rows, values):
        for index, (cell, value) in enumerate(zip(row.cells, row_values)):
            cell.width = widths[index]; set_cell_margins(cell)
            shade(cell, GREEN if index % 2 == 0 else GREEN_SOFT)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value); set_font(run, 9.5 if index % 2 == 0 else 10.5, True, "FFFFFF" if index % 2 == 0 else GREEN, "微软雅黑")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_note(doc, text, risk=False):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0); shade(cell, GOLD_SOFT if risk else GREEN_SOFT); set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]; add_rich_text(p, text, 9.5, "7D430F" if risk else GREEN)


def add_markdown_table(doc, rows):
    columns = len(rows[0])
    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = True
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]; shade(cell, GREEN); set_cell_margins(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, value, 7.0 if columns >= 7 else 8.5, "FFFFFF")
        for run in p.runs: run.bold = True
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]; set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2: shade(cell, "F5F7F5")
            if index == 0 and ("学院" in rows[0][0] or "中心" in rows[0][0]): shade(cell, GREEN_SOFT)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, value, 6.9 if columns >= 7 else 8.3)
            if index in (0, 1, 2):
                for run in p.runs: run.bold = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table_caption(doc, rows, number):
    header = "|".join(rows[0])
    if "招生学院" in header and "2027拟招人数" in header:
        title = "大连海事大学2027年电子通信相关专业与初复试科目"
        scope = "涉及单位：信息科学技术学院、港口与航运安全协同创新中心、水路交通控制全国重点实验室"
    elif "2024复试线" in header:
        title = "大连海事大学2024—2026年电子通信方向复试线"
        scope = "学院：信息科学技术学院"
    elif "复试人数" in header and "录取平均分" in header:
        title = "大连海事大学2026年一志愿复试录取情况"
        scope = "学院：信息科学技术学院"
    elif "专业课最高分" in header:
        title = "大连海事大学2026年807信号与系统成绩统计"
        scope = "学院：信息科学技术学院"
    elif "调剂录取最高分" in header:
        title = "大连海事大学2026年电子通信相关专业调剂录取情况"
        scope = "按学院、中心及重点实验室分别统计"
    elif "学术型硕士" in header and "专业型硕士" in header:
        title = "大连海事大学2027年学费、奖助学金与住宿信息"
        scope = "口径：学术型硕士与专业型硕士"
    elif "考生背景" in header:
        title = "大连海事大学电子通信方向报考建议"
        scope = "依据：招生规模、录取分数、专业课表现与数据风险"
    else:
        title = "大连海事大学电子通信考研数据表"
        scope = "数据来源：2027版择校手册"
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"表{number} {title}"); set_font(r, 10.5, True, INK, "微软雅黑")
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(scope); set_font(r, 8.5, False, MUTED, "微软雅黑")
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True


def build(output: Path = OUTPUT, source: Path = SOURCE, chart: Path = CHART):
    lines = Path(source).read_text(encoding="utf-8").splitlines()
    doc = Document(); style_document(doc)
    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(style="Title"); add_rich_text(p, title, 22)
    p = doc.add_paragraph(); r = p.add_run("2027电子通信考研 · 单校择校分析"); set_font(r, 10, True, GOLD, "微软雅黑")
    add_identity_band(doc)
    table_buffer = []
    chart_added = False
    current_heading = ""
    admission_college_added = False
    table_number = 0
    index = 1
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("| "):
            table_buffer = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [item.strip() for item in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", item) for item in cells): table_buffer.append(cells)
                index += 1
            if table_buffer:
                table_number += 1
                add_table_caption(doc, table_buffer, table_number)
                add_markdown_table(doc, table_buffer)
            continue
        if line.startswith("[图表建议"):
            if "复试线变化" in line and not chart_added and Path(chart).exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(chart), width=Cm(16.0))
                caption = doc.add_paragraph("图1 大连海事大学信息科学技术学院2024—2026年电子通信相关专业复试线变化（单位：分）")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in caption.runs: set_font(run, 8.5, False, MUTED)
                chart_added = True
            index += 1; continue
        if not line:
            index += 1; continue
        if line.startswith("# "):
            pass
        elif line.startswith("## "):
            source_heading = line[3:]
            heading_map = {
                "一、大连海事大学值不值得关注？": "院校定位",
                "二、四个电子通信方向有什么区别？": "一、2027初试科目",
                "三、近三年复试线：波动比想象中更大": "二、2026复试线",
                "四、2026录取难度：复试线不是最终门槛": "三、2026录取情况",
                "五、807信号与系统难不难？": "四、专业课难度",
                "六、调剂、复试、学费与住宿": "五、调剂、复试与就读信息",
                "七、不同分数和背景的考生怎么选？": "六、报考难度分析",
            }
            current_heading = heading_map.get(source_heading, source_heading)
            p = doc.add_paragraph(current_heading, style="Heading 1")
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2")
        elif re.match(r"^\d+\.\s+\*\*", line):
            item = re.sub(r"^\d+\.\s+", "", line)
            match = re.match(r"\*\*(\d{6})\s+([^*：]+)(?:：)?\*\*(.*)", item)
            if match and current_heading == "三、2026录取情况":
                if not admission_college_added:
                    doc.add_paragraph("信息科学技术学院", style="Heading 2")
                    admission_college_added = True
                p = doc.add_paragraph(f"{match.group(1)} {match.group(2).strip()}", style="Heading 3")
                detail = match.group(3).lstrip("：: ")
                if detail:
                    p = doc.add_paragraph(); add_rich_text(p, detail)
            else:
                p = doc.add_paragraph(style="List Bullet"); add_rich_text(p, item)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_rich_text(p, line[2:])
        elif line.startswith("*注") or "数学上不能同时成立" in line:
            add_note(doc, line.replace("*", ""), risk=True)
        elif line.startswith("**结论："):
            add_note(doc, line)
        else:
            p = doc.add_paragraph(); add_rich_text(p, line)
        index += 1
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("研途编辑 · 数据来源与风险说明见文末")
    set_font(run, 8.5, False, MUTED, "微软雅黑")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build())

