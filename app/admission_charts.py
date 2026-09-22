from pathlib import Path
from copy import deepcopy
from PIL import Image, ImageDraw, ImageFont
from docx.table import Table

def chart(bands, college, code, name, target, scope_label=''):
    data=sorted(bands,key=lambda x:int(x['band'].split('-')[0]))
    im=Image.new('RGB',(1600,900),'white'); d=ImageDraw.Draw(im)
    def font(n,b=False): return ImageFont.truetype('C:/Windows/Fonts/'+('msyhbd.ttc' if b else 'msyh.ttc'),n)
    def center(x,y,text,size,b=False):
        box=d.textbbox((0,0),text,font=font(size,b));d.text((x-(box[2]-box[0])/2,y),text,font=font(size,b),fill='black')
    center(800,25,'2026年一志愿初试分数段分布',34,True)
    center(800,76,f'{college} - {code} {name}',23,True)
    left,right,top,bottom=125,1550,180,775
    highest=max(x['retest_count'] for x in data)
    step=max(1,(highest+5)//6);limit=((highest+step-1)//step)*step
    if limit==highest:limit+=step
    for value in range(0,limit+1,step):
        y=bottom-(bottom-top)*value/limit
        for x in range(left,right,13):d.line((x,y,min(x+7,right),y),fill='#dddddd',width=1)
        d.text((left-38,y-12),str(value),font=font(21),fill='black')
    d.line((left,top,left,bottom,right,bottom),fill='black',width=2)
    d.text((35,430),'人\n数',font=font(25),fill='black',spacing=5)
    for i,(label,color) in enumerate([('进入复试','#4472C4'),('拟录取','#FFC000')]):
        y=125+i*32;d.rectangle((145,y,180,y+19),fill=color);d.text((192,y-6),label,font=font(23),fill='black')
    group=(right-left)/len(data);bw=min(53,group*.34)
    for i,row in enumerate(data):
        x=left+group*(i+.5)
        for j,key in enumerate(('retest_count','admitted')):
            v=row[key];a=x-bw if j==0 else x;y=bottom-(bottom-top)*v/limit
            if v:d.rectangle((a,y,a+bw,bottom-1),fill=('#4472C4','#FFC000')[j]);center(a+bw/2,y-29,str(v),21)
        center(x,bottom+17,row['band'],20)
    center(800,827,'初试分数段',27)
    if scope_label:center(800,870,scope_label,19,True)
    target=Path(target);target.parent.mkdir(parents=True,exist_ok=True);im.save(target)

def reference_table(doc, headers, row):
    path=Path(__file__).resolve().parents[1]/'templates/admission_summary_reference.xml'
    if not path.exists():return False
    from docx.oxml import parse_xml
    node=parse_xml(path.read_bytes());doc.element.body.insert(len(doc.element.body)-1,node)
    table=Table(node,doc)
    for cells,values in ((table.rows[0].cells,headers),(table.rows[1].cells,row)):
        for cell,value in zip(cells,values):
            p=cell.paragraphs[0]
            if not p.runs:p.add_run('')
            p.runs[0].text=str(value)
            for run in p.runs[1:]:run.text=''
            for extra in cell.paragraphs[1:]:extra._p.getparent().remove(extra._p)
    return True

