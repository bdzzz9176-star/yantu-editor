from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from admission_model import normalize_facts, require_confirmed
from admission_document import FORMAT, MARKER, compile_draft, resolve_asset
from admission_assets import collect_assets, prepare_assets

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


GREEN, GREEN_SOFT, GOLD, GOLD_SOFT, INK, MUTED, LINE = "18563F", "E8F0EB", "B66A20", "FFF1DC", "1A211E", "68736F", "D8DDD9"


def assess_draft_quality(content: str) -> dict[str, Any]:
    findings=[]
    compact=re.sub(r"\s+","",content)
    if len(compact)<3500:
        findings.append({"severity":"block","code":"too_short","detail":f"正文仅约{len(compact)}字符，尚未达到可交编辑审核的完整度（最低3500字符）。"})
    table_count=len(re.findall(r"^\|.+\|$",content,re.MULTILINE))
    if table_count<6:
        findings.append({"severity":"block","code":"insufficient_tables","detail":"有效数据表不足，学校、学院与专业数据尚未形成完整分析。"})
    reader_body=content
    leaks=[x for x in ("流程验证稿","测试稿","JSON输入","内部工作流","提示词","配置API","经用户确认","按用户确认","文章主表","内部口径","程序保留","不扣减","核心结论先行","所有判断仅依据已确认","手册覆盖","择校手册","手册未披露","资料未披露","发布前核验","来源页码","资料来源与说明") if x in reader_body]
    if leaks:
        findings.append({"severity":"block","code":"internal_process_leak","detail":leaks})
    headings=len(re.findall(r"^##\s+",content,re.MULTILINE))
    if headings<5:
        findings.append({"severity":"block","code":"incomplete_structure","detail":"一级分析章节不足5个，文章结构不完整。"})
    required=("学校简介与学科实力","考试科目及参考书目","近三年复试线","2026拟录取分析","复试内容","总体难度分析")
    missing=[x for x in required if not re.search(rf'^##\s+{re.escape(x)}\s*$',content,re.MULTILINE)]
    if not re.search(r'^##\s+专业课难度(?:分析（复试名单未公布专业课分数）)?\s*$',content,re.MULTILINE):missing.append('专业课难度')
    if not re.search(r'^##\s+20\d{2}初复试科目\s*$',content,re.MULTILINE):missing.append('目标年份初复试科目')
    matched_template_headings=sum(f'## {name}' in content for name in required)
    template_draft=matched_template_headings>=3 or bool(re.search(r'^##\s+20\d{2}初复试科目',content,re.MULTILINE))
    if template_draft and missing:findings.append({"severity":"block","code":"template_structure_mismatch","detail":"缺少范文章节："+'、'.join(missing)})
    unsupported=[]
    checks=[r'目标分数建议',r'初试目标建议',r'建议定在\s*\d+\s*分',r'录取概率(?:明显|较|偏|将)',r'不排除专项',r'可能(?:属于|涉及|来自)专项',r'专项考生可能']
    for pattern in checks:
        if re.search(pattern,content):unsupported.append(pattern)
    if unsupported:
        findings.append({"severity":"block","code":"unsupported_admission_inference","detail":"存在手册不能证明的目标分或专项/概率推断："+'、'.join(unsupported)})
    placeholders=[x for x in ('None','现有数据未包含','暂无相关数据') if x in content]
    if placeholders:
        findings.append({"severity":"block","code":"reader_placeholder_leak","detail":"正文仍有面向内部的数据占位语："+'、'.join(placeholders)})
    return {"generation_ready_for_word":not any(x["severity"]=="block" for x in findings),"findings":findings,"metrics":{"characters":len(compact),"table_rows":table_count,"sections":headings}}


def task_dir(workspace: Path, task_id: str) -> Path:
    return workspace / "tasks" / task_id


def load_bundle(workspace: Path, tasks_file: Path, task_id: str) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    tasks = json.loads(tasks_file.read_text(encoding="utf-8"))
    task = next((item for item in tasks if item.get("task_id") == task_id), None)
    if not task:
        raise ValueError("任务不存在")
    root = task_dir(workspace, task_id)
    facts_path, outline_path = root / "structured_facts.json", root / "outline.json"
    if not facts_path.is_file() or not outline_path.is_file():
        raise ValueError("请先完成数据与框架确认")
    facts=normalize_facts(json.loads(facts_path.read_text(encoding="utf-8")))
    if facts.get('school')!=task.get('school'):raise ValueError('任务与事实库学校不一致')
    return task, facts, json.loads(outline_path.read_text(encoding="utf-8"))


def task_assets(workspace,task_id,facts):
    path=task_dir(workspace,task_id)/'source_extract.json'
    if not path.is_file():raise ValueError('缺少当前任务原始抽取，无法核验分段')
    return collect_assets(facts,json.loads(path.read_text(encoding='utf-8'))['pages'])


def read_task_draft(workspace,task_id,facts):
    out=task_dir(workspace,task_id)/'outputs'
    meta=json.loads((out/'draft_meta.json').read_text(encoding='utf-8')) if (out/'draft_meta.json').is_file() else {}
    content=(out/'draft.md').read_text(encoding='utf-8')
    content=compile_draft(content,facts,task_assets(workspace,task_id,facts),add_default_charts=meta.get('draft_format')!=FORMAT and not MARKER.search(content))
    content=content.split('\n## 资料来源与说明',1)[0].rstrip()+'\n'
    content=content.replace('择校手册','资料').replace('手册未披露','').replace('资料未披露','')
    # Sections consisting only of a missing-data notice do not belong in the article.
    content=re.sub(r'\n## (?:调剂情况|往年录取情况)\s*\n+(?:现有数据|资料|当前资料)[^\n]*(?:未包含|暂无)[^\n]*\n?', '\n', content)
    return content


def build_prompt(workspace: Path, tasks_file: Path, project_root: Path, task_id: str) -> tuple[str, str]:
    task, facts, outline = load_bundle(workspace, tasks_file, task_id)
    require_confirmed(task)
    if any('待确认' in str(r.get('college','')) for r in facts.get('admission_2026',[])):
        raise ValueError('录取数据学院归属未确认，不能生成正式稿')
    if outline.get("status") != "approved" or not outline.get("selected_title_id"):
        raise ValueError("文章框架尚未确认")
    selected = next(item for item in outline["title_options"] if item["id"] == outline["selected_title_id"])
    system = (project_root / "prompts" / "single_school_system.md").read_text(encoding="utf-8")
    system += "\n\n当前是任务隔离模式。只能使用SOURCE_FACTS中的事实和数字；不得引用其他学校案例。参考范文只决定写法，不提供事实。表格标题必须写清学校、学院、年份和统计口径。"
    payload = {"task": {"school": task.get("school"), "target_year": task.get("target_year")}, "confirmed_title": selected["title"], "audience": outline["audience"], "opening_angle": outline["opening_angle"], "sections": outline["sections"], "tables": outline["tables"], "charts": outline["charts"], "required_disclosures": outline["required_disclosures"], "user_notes": (outline.get("approval") or {}).get("notes", ""), "SOURCE_FACTS": facts}
    assets=task_assets(workspace,task_id,facts)
    band_assets=[asset for asset in assets if asset.get('bands')]
    payload['charts']=outline['charts'] if band_assets else []
    payload['verified_admission_bands']=[{k:v for k,v in asset.items() if k!='scope_note'} for asset in band_assets]
    payload['admission_sections']=[{'key':a['key'],'college':a['college'],'code':a['code'],'name':a['name'],'placeholder':'[录取分析:'+a['key']+']'} for a in assets]
    from handbook_tables import standard_placeholders
    payload['handbook_sections']=standard_placeholders(facts)
    system += '\n成稿采用所给四川大学范文的章节名称、顺序和表达标准。初复试科目章节必须命名为“2027初复试科目”；“考试科目及参考书目”紧随其后，集中放初试专业课和复试参考书。禁止输出“核心结论先行”标题、提纲、项目符号结论或编辑说明；导语只写一个自然段。录取分析按学院、专业逐项展开，在各专业标题后原样放入admission_sections中的placeholder，随后最多写两段短分析：第一段概括人数、均分和最低分，第二段只概括最有解释力的一两个分段，不得逐档枚举全部分数段。报考建议只写简短自然段，不得给目标分、稳妥区间、录取概率。程序会展开标准横向表，不得另写纵向重复表，也不得输出[图表建议]。全文禁止出现“择校手册、手册、资料披露、资料未披露、来源页码、用户确认、发布前核验”等资料整理语言；缺少的字段直接不写，不得编造。'
    system += '\nSOURCE_FACTS中已识别的通用栏目由handbook_sections提供。必须在对应章节逐个原样放入其中的placeholder，程序会展开为带依据的标准表；不得手抄、删列或省略。'
    system += '\n术语必须区分：admit_rate是录取率（录取人数÷复试人数），复录比是复试人数÷录取人数，不能混写。不得把低于普通复试线的样本推断为专项，不得推测专项所在分数段；不得把专业课均分换算成正确率；不得给出模型自行计算的目标分、稳妥线或录取概率。只比较事实库直接提供的分数、人数和比例，并明确样本与统计口径边界。'
    system += '\nSOURCE_FACTS中的scope_note是内部审计记录，禁止原样或改写后输出。不得在正文出现“用户确认、文章主表、内部口径、程序保留、不扣减、不推测”等编辑流程语言。统计范围差异只使用reader_scope_note中的读者说明，并放在文末数据说明。'
    return system, "请严格依据以下JSON生成完整公众号初稿：\n\n" + json.dumps(payload, ensure_ascii=False, indent=2)


def persist_draft(workspace: Path, task_id: str, content: str, usage: dict[str, Any], operation: str = "generation", instructions: str = "", compliance: list[dict] | None = None) -> dict[str, Any]:
    from source_notes import sources_at_end
    facts=normalize_facts(json.loads((task_dir(workspace,task_id)/'structured_facts.json').read_text(encoding='utf-8')))
    content=compile_draft(content,facts,task_assets(workspace,task_id,facts),add_default_charts=operation=='generation',require_handbook_sections=operation=='generation')
    out = task_dir(workspace, task_id) / "outputs"; out.mkdir(parents=True, exist_ok=True)
    draft = out / "draft.md"
    if draft.exists():
        history = out / "draft_history"; history.mkdir(exist_ok=True)
        from shutil import copy2
        copy2(draft,history / f"draft_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.md")
    draft.write_text(content.rstrip() + "\n", encoding="utf-8")
    meta = {"generated_at": datetime.now(timezone.utc).isoformat(), "operation": operation, "usage": usage, "draft_format":FORMAT}
    (out / "draft_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    versions_file=out/'versions.json'; versions=json.loads(versions_file.read_text(encoding='utf-8')) if versions_file.is_file() else []
    version_id=len(versions)+1; version_dir=out/'versions'; version_dir.mkdir(exist_ok=True)
    snapshot=version_dir/f'draft_v{version_id}.md'; snapshot.write_text(content.rstrip()+"\n",encoding='utf-8')
    versions.append({'version_id':version_id,'created_at':meta['generated_at'],'operation':operation,'instructions':instructions,'compliance':compliance or [],'draft_file':str(snapshot.relative_to(out))})
    versions_file.write_text(json.dumps(versions,ensure_ascii=False,indent=2)+"\n",encoding='utf-8')
    meta['version_id']=version_id
    return meta


def build_validation_draft(workspace: Path, tasks_file: Path, task_id: str) -> str:
    """Build a factual local draft for end-to-end validation when the API key is unavailable."""
    task,facts,outline=load_bundle(workspace,tasks_file,task_id)
    selected=next(x for x in outline["title_options"] if x["id"]==outline["selected_title_id"])
    programs=facts.get("programs",[]); admissions=facts.get("admission_2026",[]); cutoffs=facts.get("cutoffs",[])
    names={x["code"]:x.get("name","") for x in programs}
    lines=[f"# {selected['title']}","",outline["opening_angle"],""]
    lines += ["## 一、院校与招生范围", "", f"从手册字段看，{task['school']}电子通信相关招生共识别出{len(programs)}组学院与专业组合。择校时不能只看学校名称，还要把学院、专业代码、统招人数和考试科目放在一起判断。", ""]
    lines += ["## 二、招生专业与初复试科目","","| 学院 | 专业 | 统招人数 | 初试科目 | 复试 | 来源 |","|---|---|---:|---|---|---|"]
    for x in programs: lines.append(f"| {x.get('college',facts.get('college',''))} | {x['code']} {x['name']} | {x.get('planned','待确认')} | {x.get('initial_subjects','待确认')} | {x.get('retest','待确认')} | 手册第{x['page']}页 |")
    lines += ["", "不同专业即使代码相近，也可能属于不同学院并使用不同专业课。统招人数较少的方向受个别考生分数影响更明显，不能仅凭某一年的最低分判断稳定性。", ""]
    if cutoffs:
        lines += ["## 三、近三年复试线","","| 专业 | 2024 | 2025 | 2026 | 变化 | 来源 |","|---|---:|---:|---:|---:|---|"]
        for x in cutoffs: lines.append(f"| {x.get('code') or x.get('name','')} | {x['2024']} | {x['2025']} | {x['2026']} | {x.get('change','')} | 手册第{x['page']}页 |")
        lines += ["", "复试线反映的是进入复试的门槛，不等于最终录取的普遍分数。连续三年的变化更适合观察方向热度和波动，单年低点不应被包装成容易上岸。", ""]
    if admissions:
        lines += ["## 四、2026年一志愿录取情况","","| 学院/专业 | 最高分 | 最低分 | 中位分 | 平均分 | 复试人数 | 录取人数 | 录取率 | 来源 |","|---|---:|---:|---:|---:|---:|---:|---:|---|"]
        for x in admissions: lines.append(f"| {x.get('college','')} {x['code']} {x.get('name',names.get(x['code'],''))} | {x['max']} | {x['min']} | {x['median']} | {x['average']} | {x['retest_count']} | {x['admitted']} | {x['admit_rate']} | 手册第{x['page']}页 |")
        low=min(admissions,key=lambda x:x['min']); avg=max(admissions,key=lambda x:x['average'])
        lines += ["",f"最低录取分较低的是{low.get('college','')}{low['code']}，为{low['min']}分；但这只是样本下沿。平均分较高的是{avg.get('college','')}{avg['code']}，为{avg['average']}分。判断真实难度时，应同时参考录取平均分、复试人数和录取率。",""]
        lines += ["## 五、专业课与备考要求","","| 专业 | 专业课最高 | 专业课平均 | 专业课最低 | 数学平均 | 来源 |","|---|---:|---:|---:|---:|---|"]
        for x in admissions: lines.append(f"| {x['code']} {x.get('name',names.get(x['code'],''))} | {x.get('professional_max','待确认')} | {x.get('professional_average','待确认')} | {x.get('professional_min','待确认')} | {x.get('math_average','待确认')} | 手册第{x['page']}页 |")
        lines += ["", "专业课平均分可以辅助判断试题与考生基础，但不能脱离科目组成和样本规模单独比较。备考前应先确认目标学院的专业课代码，再按对应科目准备。", ""]
    lines += ["## 六、报考建议","", "第一，先按学院锁定专业代码，避免只看专业名称导致科目准备错误。第二，目标分数应参考录取平均分和中位分，不能只贴着复试线制定。第三，招生人数较少的方向要预留更大的分数余量。第四，发布前再次核验正式招生目录、推免和专项口径。", "", "## 数据说明", ""]
    lines += [f"- {x}" for x in outline["required_disclosures"]]
    lines += ["", "> 当前版本为结构化数据生成的流程验证稿，用于检查数据、图表和Word链路；配置模型后可在页面重新生成更完整的正式初稿。"]
    return "\n".join(lines)+"\n"


def generate_charts(workspace: Path, tasks_file: Path, task_id: str) -> dict[str, Any]:
    task,facts,_=load_bundle(workspace,tasks_file,task_id)
    require_confirmed(task)
    draft=task_dir(workspace,task_id)/'outputs/draft.md'
    if not draft.is_file():raise ValueError('请先生成并确认正文中的图表位置')
    content=read_task_draft(workspace,task_id,facts)
    requested={m[1] for m in MARKER.finditer(content)}
    return prepare_assets(workspace,task_id,facts,requested_keys=requested)


def _set_font(run, size=10.5, bold=False, color=INK, name="微软雅黑"):
    run.font.name=name; fonts=run._element.get_or_add_rPr().rFonts
    for key in ("w:eastAsia","w:ascii","w:hAnsi","w:cs"): fonts.set(qn(key),name)
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)


def _shade(cell, fill):
    pr=cell._tc.get_or_add_tcPr(); node=pr.find(qn("w:shd")) or OxmlElement("w:shd"); node.set(qn("w:fill"),fill)
    if node.getparent() is None: pr.append(node)


def _table(doc, headers, rows, widths):
    table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    for i,(cell,text) in enumerate(zip(table.rows[0].cells,headers)):
        table.columns[i].width=Cm(widths[i])
        cell.width=Cm(widths[i]); _shade(cell,"DDEBF7"); p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(str(text)); _set_font(r,8,True,"000000")
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; _set_font(p.add_run(str(text)),7.6)
    return table


def build_word(workspace: Path, tasks_file: Path, task_id: str) -> Path:
    task,facts,outline=load_bundle(workspace,tasks_file,task_id); out=task_dir(workspace,task_id)/"outputs"; draft=out/"draft.md"
    require_confirmed(task)
    if outline.get('status')!='approved' or not outline.get('selected_title_id'):
        raise ValueError('文章框架尚未确认')
    if not draft.is_file(): raise ValueError("请先生成正文初稿")
    meta_path=out/"draft_meta.json"
    meta=json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.is_file() else {}
    if meta.get("operation")=="validation_fallback":
        raise ValueError("当前只是流程预览稿，未达到正式稿标准，不能导出Word。请先调用模型生成正式初稿。")
    content=read_task_draft(workspace,task_id,facts)
    quality=assess_draft_quality(content)
    if not quality["generation_ready_for_word"]:
        raise ValueError("当前初稿未通过正式稿质量检查，不能导出Word："+"；".join(str(x["detail"]) for x in quality["findings"]))
    requested={m[1] for m in MARKER.finditer(content)}
    inline_assets=prepare_assets(workspace,task_id,facts,requested_keys=requested)['charts']
    doc=Document(); sec=doc.sections[0]; sec.page_width,sec.page_height=Cm(21),Cm(29.7); sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
    normal=doc.styles["Normal"]; normal.font.name="微软雅黑"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑"); normal._element.rPr.rFonts.set(qn("w:ascii"),"微软雅黑"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"微软雅黑"); normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.3; normal.paragraph_format.space_after=Pt(4)
    for name,size,color,before,after in [("Title",22,INK,0,12),("Heading 1",15,GREEN,14,6),("Heading 2",12.5,INK,10,5)]:
        s=doc.styles[name]; s.font.name="微软雅黑"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    title=next(x["title"] for x in outline["title_options"] if x["id"]==outline["selected_title_id"]); p=doc.add_paragraph(title,style="Title"); doc.add_paragraph(f"{task.get('target_year','2027')}电子通信考研 · 单校择校分析")
    colleges=facts.get("colleges") or [{"name":facts.get("college","待确认"),"code":facts.get("college_code","")}]; p=doc.add_paragraph(); _set_font(p.add_run(f"学校：{task['school']}　学院："+"、".join(f"{x['name']}（{x['code']}）" for x in colleges)),10,True,GREEN,"微软雅黑")
    from source_notes import sources_at_end
    lines=content.splitlines()[1:]
    index=0
    while index < len(lines):
        text=lines[index].strip()
        if not text: index+=1; continue
        if text.startswith("|") and index+1<len(lines) and re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*",lines[index+1]):
            block=[]
            while index<len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip()); index+=1
            parsed=[[c.strip() for c in row.strip("|").split("|")] for row in block]
            headers=parsed[0]; rows=parsed[2:]
            if rows and headers:
                from admission_charts import reference_table
                if headers[0]=='学院及专业' and len(headers)==7 and len(rows)==1 and reference_table(doc,headers,rows[0]):
                    continue
                weights=[3 if '初试科目' in h else 2 if any(k in h for k in ('学院','名称','复试科目')) else 1 for h in headers]
                _table(doc,headers,rows,[15.6*w/sum(weights) for w in weights])
            continue
        if re.fullmatch(r"[:| -]+",text): index+=1; continue
        if text.startswith('[录取分段图:'):
            token=text.split(':',1)[1].rstrip(']')
            asset=resolve_asset(token,inline_assets)
            target=out/'charts'/asset['filename']
            if not target.is_file():raise ValueError('缺少已核验的分数段图，请先生成图表')
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(target),width=Cm(15.6))
        elif text=='## 2026拟录取分析':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            _set_font(p.add_run('2026拟录取分析'),14,True,'000000')
        elif text.startswith("## "): doc.add_paragraph(text[3:],style="Heading 1")
        elif text.startswith("### "): doc.add_paragraph(text[4:],style="Heading 2")
        elif text.startswith("- "): doc.add_paragraph(text[2:],style="List Bullet")
        elif text.startswith("**") and text.endswith("**"):
            p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            _set_font(p.add_run(text[2:-2]),11,True,GREEN,"微软雅黑")
        elif not text.startswith("[图表建议"): doc.add_paragraph(re.sub(r"\*\*([^*]+)\*\*",r"\1",text))
        index+=1
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; _set_font(footer.add_run("研途编辑"),8.5,False,MUTED)
    # Enforce one font across every generated run, including inherited table and heading runs.
    for paragraph in list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]:
        for run in paragraph.runs:
            rpr=run._element.get_or_add_rPr(); fonts=rpr.rFonts
            if fonts is None:
                fonts=OxmlElement("w:rFonts"); rpr.insert(0,fonts)
            for key in ("w:eastAsia","w:ascii","w:hAnsi","w:cs"): fonts.set(qn(key),"微软雅黑")
            run.font.name="微软雅黑"
    latest=out/f"{task['school']}电子通信考研择校分析_最新.docx"; history=out/"word_history"
    if latest.exists():
        history.mkdir(exist_ok=True)
        try: latest.replace(history/f"{task['school']}电子通信考研择校分析_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        except PermissionError:
            unlocked=out/f"{task['school']}电子通信考研择校分析_模板修正版.docx"; doc.save(unlocked); return unlocked
    doc.save(latest); return latest

