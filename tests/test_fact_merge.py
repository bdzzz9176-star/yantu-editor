import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
from fact_merge import merge_fact_sets,preserve_reviewed_overrides
from material_extract import extract_docx_school_section
from docx import Document


class FactMergeTest(unittest.TestCase):
    def test_same_rows_are_deduplicated(self):
        row={'college_code':'001','college':'电子学院','code':'085400','name':'电子信息','planned':10}
        merged=merge_fact_sets('甲大学',[('old',{'programs':[row]}),('new',{'programs':[row]})])
        self.assertEqual(len(merged['programs']),1)
        self.assertEqual(merged['programs'][0]['source_material_ids'],['old','new'])
        self.assertEqual(merged['source_conflicts'],[])

    def test_later_source_wins_and_conflict_is_visible(self):
        old={'college_code':'001','college':'电子学院','code':'085400','name':'电子信息','planned':10}
        new={**old,'planned':12}
        merged=merge_fact_sets('甲大学',[('old',{'programs':[old]}),('new',{'programs':[new]})])
        self.assertEqual(merged['programs'][0]['planned'],12)
        self.assertEqual(merged['source_conflicts'][0]['values'],[10,12])
        self.assertIn('10 → 12',merged['warnings'][-1])

    def test_word_handbook_extracts_only_target_school_and_keeps_table_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'handbook.docx';doc=Document()
            doc.add_paragraph('甲大学');doc.add_paragraph('一、院校基本信息');doc.add_paragraph('甲校内容')
            table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='录取平均分';table.cell(0,1).text='382.5'
            doc.add_paragraph('乙大学');doc.add_paragraph('一、院校基本信息');doc.add_paragraph('乙校内容');doc.save(path)
            text=extract_docx_school_section(path,'甲大学')
            self.assertIn('录取平均分\n382.5',text)
            self.assertNotIn('乙校内容',text)

    def test_reparse_preserves_reviewed_reporting_scope(self):
        identity={'college_code':'205','code':'085400','name':'电子信息'}
        old={'admission_2026':[{**identity,'reporting':{'admitted':83},'reader_scope_note':'普通统招与含专项口径不同'}]}
        new={'admission_2026':[{**identity,'admitted':93}]}
        result=preserve_reviewed_overrides(old,new)
        self.assertEqual(result['admission_2026'][0]['reporting']['admitted'],83)
        self.assertIn('普通统招',result['admission_2026'][0]['reader_scope_note'])


if __name__=='__main__':unittest.main()

